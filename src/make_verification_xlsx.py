"""
make_verification_xlsx.py - builds results/IrrigAI_member3_excel_verification.xlsx

Purpose: let Member 3 (and the viva examiner) repeat the Python numbers with plain
Excel formulas (COUNTIFS, SUMPRODUCT, CHIDIST). The Excel columns are LIVE formulas on the
`data` sheet; the "Python" columns are typed-in values copied from results/*.csv, so each row
can be compared with a MATCH / CHECK flag.

Run (after irrigai_analysis_g6.py):  python src/make_verification_xlsx.py
Then recalculate so the formula results are stored (LibreOffice does this in the audit).
"""
from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

ROOT = Path(__file__).resolve().parents[1]
R = ROOT / "results"
df = pd.read_csv(ROOT / "data" / "g6_irrigation_data.csv", keep_default_na=False, na_values=[""])
N = len(df)

ARIAL = Font(name="Arial", size=10)
BOLD = Font(name="Arial", size=10, bold=True)
BLUE = Font(name="Arial", size=10, color="0000FF")      # typed-in Python values (inputs)
HEAD = PatternFill("solid", fgColor="DDEBF7")
wb = Workbook()

# ---------------- data sheet ----------------
ws = wb.active
ws.title = "data"
ws.append(list(df.columns))
for row in df.itertuples(index=False):
    ws.append(list(row))
for c in ws[1]:
    c.font, c.fill = BOLD, HEAD
last = N + 1                                   # last data row on the sheet
rng = lambda col: f"data!${col}$2:${col}${last}"   # noqa: E731
RAIN, SOIL, NEXT = rng("D"), rng("E"), rng("F")

# ---------------- worked 20-row sheet ----------------
w = wb.create_sheet("worked_20_rows")
w20 = pd.read_csv(R / "worked_example_20rows.csv", keep_default_na=False, na_values=[""])
cnt = pd.read_csv(R / "worked_example_20rows_counts.csv").set_index("quantity")["value"]
w["A1"] = "Worked example: 20 consecutive real rows (demonstration only). Formulas below count them in Excel."
w["A1"].font = BOLD
w.append([])
w.append(list(w20.columns))
for c in w[3]:
    c.font, c.fill = BOLD, HEAD
for row in w20.itertuples(index=False):
    w.append(list(row))
first, lastw = 4, 3 + len(w20)
S20, R20, Y20 = (f"$E${first}:$E${lastw}", f"$D${first}:$D${lastw}", f"$F${first}:$F${lastw}")
r0 = lastw + 2
w.cell(r0, 1, "Quantity").font = BOLD
w.cell(r0, 2, "Excel formula result").font = BOLD
w.cell(r0, 3, "Python value (typed)").font = BOLD
w.cell(r0, 4, "Match?").font = BOLD
items = [
    ("N (rows)", f"=COUNTA($A${first}:$A${lastw})", cnt["N"]),
    ("n(Dry)", f'=COUNTIFS({S20},"Dry")', cnt["n(Dry)"]),
    ("n(None)", f'=COUNTIFS({R20},"None")', cnt["n(None)"]),
    ("n(Dry and None)", f'=COUNTIFS({S20},"Dry",{R20},"None")', cnt["n(Dry and None)"]),
    ("n(Dry and next Dry)", f'=COUNTIFS({S20},"Dry",{Y20},1)', cnt["n(Dry and next Dry)"]),
    ("n(Dry and Light)", f'=COUNTIFS({S20},"Dry",{R20},"Light")', cnt["n(Dry and Light)"]),
    ("n(Dry and Light and next Dry)", f'=COUNTIFS({S20},"Dry",{R20},"Light",{Y20},1)', cnt["n(Dry and Light and next Dry)"]),
]
for i, (lab, f, pyv) in enumerate(items, start=r0 + 1):
    w.cell(i, 1, lab).font = ARIAL
    w.cell(i, 2, f).font = ARIAL
    w.cell(i, 3, int(pyv)).font = BLUE
    w.cell(i, 4, f'=IF(B{i}=C{i},"MATCH","CHECK")').font = ARIAL
i = r0 + len(items) + 1
nrow, ndry, nnone, ndn, kdry = (r0 + 1, r0 + 2, r0 + 3, r0 + 4, r0 + 5)
w.cell(i, 1, "P(Dry) = n(Dry)/N").font = ARIAL
w.cell(i, 2, f"=B{ndry}/B{nrow}").font = ARIAL
w.cell(i + 1, 1, "P(None) = n(None)/N").font = ARIAL
w.cell(i + 1, 2, f"=B{nnone}/B{nrow}").font = ARIAL
w.cell(i + 2, 1, "P(Dry and None) = n(Dry and None)/N").font = ARIAL
w.cell(i + 2, 2, f"=B{ndn}/B{nrow}").font = ARIAL
w.cell(i + 3, 1, "P(Dry) x P(None)").font = ARIAL
w.cell(i + 3, 2, f"=B{i}*B{i+1}").font = ARIAL
w.cell(i + 4, 1, "P(soil_next_dry=1 | Dry) = n(Dry and next Dry)/n(Dry)").font = ARIAL
w.cell(i + 4, 2, f"=B{kdry}/B{ndry}").font = ARIAL
for rr in range(i, i + 5):
    w.cell(rr, 2).number_format = "0.0000"

# ---------------- verification sheet ----------------
v = wb.create_sheet("verification", 1)
v["A1"] = "Excel cross-check of Python results (full dataset). Black = Excel formulas, blue = values typed in from results/*.csv"
v["A1"].font = BOLD
hdr = ["Scenario", "soil_cat", "rain_cat", "n (Excel)", "k next-day Dry (Excel)", "p (Excel)",
       "n (Python)", "k (Python)", "p (Python)", "Match?"]
for j, h in enumerate(hdr, start=1):
    c = v.cell(3, j, h)
    c.font, c.fill, c.alignment = BOLD, HEAD, Alignment(wrap_text=True)
b9 = pd.read_csv(R / "partB_all_9_combinations.csv", keep_default_na=False, na_values=[""]).set_index(["soil_cat", "rain_cat"])
scen = [("Dry + None", "Dry", "None"), ("Moderate + Light", "Moderate", "Light"),
        ("Wet + Heavy", "Wet", "Heavy"), ("Dry + Light (only non-extreme p)", "Dry", "Light"),
        ("Dry + Heavy (n = 0, p undefined)", "Dry", "Heavy")]
for i, (lab, s_, r_) in enumerate(scen, start=4):
    v.cell(i, 1, lab); v.cell(i, 2, s_); v.cell(i, 3, r_)
    v.cell(i, 4, f'=COUNTIFS({SOIL},B{i},{RAIN},C{i})')
    v.cell(i, 5, f'=COUNTIFS({SOIL},B{i},{RAIN},C{i},{NEXT},1)')
    v.cell(i, 6, f'=IF(D{i}=0,"undefined",E{i}/D{i})')      # n = 0 -> p is undefined, not zero
    v.cell(i, 7, int(b9.loc[(s_, r_), "n_days"])).font = BLUE
    v.cell(i, 8, int(b9.loc[(s_, r_), "n_next_day_dry"])).font = BLUE
    p_py = b9.loc[(s_, r_), "p_dry_tomorrow"]
    v.cell(i, 9, "undefined" if pd.isna(p_py) else float(p_py)).font = BLUE
    v.cell(i, 10, f'=IF(D{i}=0,IF(AND(G{i}=0,H{i}=0,I{i}="undefined"),"MATCH","CHECK"),'
                  f'IF(AND(D{i}=G{i},E{i}=H{i},ABS(F{i}-I{i})<0.00005),"MATCH","CHECK"))')
    v.cell(i, 6).number_format = "0.0000"; v.cell(i, 9).number_format = "0.0000"

# observed table via COUNTIFS
v["A10"] = "Chi-square of independence in Excel"; v["A10"].font = BOLD
v["A11"] = "Observed (COUNTIFS)"; v["A11"].font = BOLD
for j, s_ in enumerate(["Dry", "Moderate", "Wet"], start=2):
    v.cell(12, j, s_).font = BOLD
v.cell(12, 5, "Row total").font = BOLD
v.cell(12, 1, "rain_cat \\ soil_cat").font = BOLD
for i, r_ in enumerate(["None", "Light", "Heavy"], start=13):
    v.cell(i, 1, r_).font = BOLD
    for j in range(2, 5):
        col = "BCD"[j - 2]
        v.cell(i, j, f'=COUNTIFS({RAIN},$A{i},{SOIL},{col}$12)')
    v.cell(i, 5, f"=SUM(B{i}:D{i})")
v.cell(16, 1, "Column total").font = BOLD
for j in range(2, 6):
    col = "BCDE"[j - 2]
    v.cell(16, j, f"=SUM({col}13:{col}15)")
v["A18"] = "Expected = row total x column total / N"; v["A18"].font = BOLD
for j, s_ in enumerate(["Dry", "Moderate", "Wet"], start=2):
    v.cell(19, j, s_).font = BOLD
for i, r_ in enumerate(["None", "Light", "Heavy"], start=20):
    v.cell(i, 1, r_).font = BOLD
    for j in range(2, 5):
        col = "BCD"[j - 2]
        v.cell(i, j, f"=$E{i-7}*{col}$16/$E$16")
        v.cell(i, j).number_format = "0.00"

t = pd.read_csv(R / "table_chi2_summary.csv").set_index("statistic")["value"]
v["A24"] = "Statistic"; v["B24"] = "Excel"; v["C24"] = "Python"; v["D24"] = "Match?"
for c in ("A24", "B24", "C24", "D24"):
    v[c].font, v[c].fill = BOLD, HEAD
v["A25"] = "Chi-square"
v["B25"] = "=SUMPRODUCT((B13:D15-B20:D22)^2/B20:D22)"
v["C25"] = float(t["chi2"])
v["D25"] = '=IF(ABS(B25-C25)<0.0001,"MATCH","CHECK")'
v["A26"] = "Degrees of freedom"
v["B26"] = "=(ROWS(B13:D15)-1)*(COLUMNS(B13:D15)-1)"
v["C26"] = int(t["dof"])
v["D26"] = '=IF(B26=C26,"MATCH","CHECK")'
v["A27"] = "p-value (CHIDIST)"
v["B27"] = "=CHIDIST(B25,B26)"
v["C27"] = float(t["p_value"])
v["D27"] = '=IF(ABS(B27-C27)/C27<0.001,"MATCH","CHECK")'
v["A28"] = "Cramer's V"
v["B28"] = "=SQRT(B25/(E16*(MIN(ROWS(B13:D15),COLUMNS(B13:D15))-1)))"
v["C28"] = float(t["cramers_v"])
v["D28"] = '=IF(ABS(B28-C28)<0.00005,"MATCH","CHECK")'
v["A29"] = "Smallest expected count"
v["B29"] = "=MIN(B20:D22)"
v["C29"] = float(t["min_expected_count"])
v["D29"] = '=IF(ABS(B29-C29)<0.0001,"MATCH","CHECK")'
for c in ("C25", "C26", "C27", "C28", "C29"):
    v[c].font = BLUE
v["B27"].number_format = "0.00E+00"; v["C27"].number_format = "0.00E+00"
v["A31"] = ("Note: the Excel p-value is for the textbook test, which treats days as independent. "
            "The circular-shift permutation check (Python only, see results/table_p_values_compared.csv) is not significant, so read the Excel p-value as optimistic.")

for sh in wb.worksheets:
    sh.column_dimensions["A"].width = 44 if sh.title != "data" else 12
    for col in "BCDEFGHIJ":
        sh.column_dimensions[col].width = 16
    for row in sh.iter_rows():
        for c in row:
            if c.font == Font():          # default font -> Arial
                c.font = ARIAL
out = R / "IrrigAI_member3_excel_verification.xlsx"
wb.save(out)
print("Saved", out)
